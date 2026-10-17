import io
import csv
from datetime import datetime
from django.utils import timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch


# ─────────────────────────────────────────────
#  OpenPyXL Excel Helpers
# ─────────────────────────────────────────────

def _get_styles():
    header_fill = PatternFill(start_color="1A2332", end_color="1A2332", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=14, bold=True, color="1A2332")
    subtitle_font = Font(name="Arial", size=9, italic=True, color="666666")
    data_font = Font(name="Arial", size=9)
    bold_data_font = Font(name="Arial", size=9, bold=True)
    alt_fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
    
    thin_border_side = Side(style='thin', color='D9D9D9')
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    return {
        'header_fill': header_fill,
        'header_font': header_font,
        'title_font': title_font,
        'subtitle_font': subtitle_font,
        'data_font': data_font,
        'bold_data_font': bold_data_font,
        'alt_fill': alt_fill,
        'border': cell_border
    }


def _autofit_columns(ws, min_width=12, max_width=50):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = cell.value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, max_width), min_width)


# ─────────────────────────────────────────────
#  Reports Exports (Excel, PDF, CSV)
# ─────────────────────────────────────────────

def generate_reports_excel(start_date, end_date, total_stock_in, total_stock_out, low_stock_list, out_of_stock_list, logs_qs):
    wb = openpyxl.Workbook()
    styles = _get_styles()

    # --- SHEET 1: Summary ---
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary["A1"] = "InvenStock — Inventory Analytics Report"
    ws_summary["A1"].font = styles['title_font']
    ws_summary["A2"] = f"Period: {start_date} to {end_date} | Generated: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws_summary["A2"].font = styles['subtitle_font']

    ws_summary.append([])
    ws_summary.append(["Key Metric", "Value", "Notes"])
    row_idx = 4
    for cell in ws_summary[row_idx]:
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.alignment = Alignment(horizontal="center" if cell.column == 2 else "left", vertical="center")

    summary_rows = [
        ("Total Stock Received (Stock In)", total_stock_in, "Units taken into inventory during selected period"),
        ("Total Stock Dispatched (Stock Out)", total_stock_out, "Units checked out / sold during selected period"),
        ("Net Stock Movement", total_stock_in - total_stock_out, "Difference between inbound and outbound units"),
        ("Low Stock Products Count", len(low_stock_list), "Products at or below their alert threshold"),
        ("Out of Stock Products Count", len(out_of_stock_list), "Products currently having zero inventory"),
    ]

    for item, val, note in summary_rows:
        row_idx += 1
        ws_summary.append([item, val, note])
        r = ws_summary[row_idx]
        r[0].font = styles['bold_data_font']
        r[1].font = styles['bold_data_font']
        r[2].font = styles['data_font']
        r[1].alignment = Alignment(horizontal="right")
        for c in r:
            c.border = styles['border']

    _autofit_columns(ws_summary)

    # --- SHEET 2: Low Stock Alerts ---
    ws_low = wb.create_sheet(title="Low Stock Items")
    ws_low.views.sheetView[0].showGridLines = True
    ws_low["A1"] = "Low Stock Alerts"
    ws_low["A1"].font = styles['title_font']
    ws_low["A2"] = f"Products requiring replenishment as of {timezone.now().strftime('%Y-%m-%d')}"
    ws_low["A2"].font = styles['subtitle_font']

    ws_low.append([])
    headers_low = ["SKU", "Product Name", "Category", "Supplier", "Current Stock", "Min Alert Level", "Deficit Units"]
    ws_low.append(headers_low)
    for cell in ws_low[4]:
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 4
    for p in low_stock_list:
        row_idx += 1
        tot = p.total_stock
        deficit = max(0, p.low_stock_alert - tot)
        cat_name = p.category.name if p.category else "—"
        sup_name = p.supplier.name if p.supplier else "—"
        ws_low.append([p.sku, p.name, cat_name, sup_name, tot, p.low_stock_alert, deficit])
        r = ws_low[row_idx]
        for c in r:
            c.font = styles['data_font']
            c.border = styles['border']
        r[4].alignment = Alignment(horizontal="right")
        r[5].alignment = Alignment(horizontal="right")
        r[6].alignment = Alignment(horizontal="right")
        if row_idx % 2 == 0:
            for c in r:
                c.fill = styles['alt_fill']

    _autofit_columns(ws_low)

    # --- SHEET 3: Out of Stock ---
    ws_out = wb.create_sheet(title="Out of Stock")
    ws_out.views.sheetView[0].showGridLines = True
    ws_out["A1"] = "Out of Stock Products"
    ws_out["A1"].font = styles['title_font']
    ws_out["A2"] = f"Zero-quantity products as of {timezone.now().strftime('%Y-%m-%d')}"
    ws_out["A2"].font = styles['subtitle_font']

    ws_out.append([])
    headers_out = ["SKU", "Product Name", "Category", "Supplier", "Retail Price", "Cost Price", "Status"]
    ws_out.append(headers_out)
    for cell in ws_out[4]:
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 4
    for p in out_of_stock_list:
        row_idx += 1
        cat_name = p.category.name if p.category else "—"
        sup_name = p.supplier.name if p.supplier else "—"
        ws_out.append([p.sku, p.name, cat_name, sup_name, float(p.retail_price), float(p.cost_price), p.get_status_display()])
        r = ws_out[row_idx]
        for c in r:
            c.font = styles['data_font']
            c.border = styles['border']
        r[4].number_format = '$#,##0.00'
        r[5].number_format = '$#,##0.00'
        if row_idx % 2 == 0:
            for c in r:
                c.fill = styles['alt_fill']

    _autofit_columns(ws_out)

    # --- SHEET 4: Stock Transactions Audit ---
    ws_logs = wb.create_sheet(title="Stock Transactions")
    ws_logs.views.sheetView[0].showGridLines = True
    ws_logs["A1"] = "Inventory Movements Audit Log"
    ws_logs["A1"].font = styles['title_font']
    ws_logs["A2"] = f"Transactions recorded between {start_date} and {end_date}"
    ws_logs["A2"].font = styles['subtitle_font']

    ws_logs.append([])
    headers_logs = ["Date / Time", "Reference", "Type", "Product SKU", "Product Name", "Qty Changed", "From Location", "To Location", "Performed By", "Reason"]
    ws_logs.append(headers_logs)
    for cell in ws_logs[4]:
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 4
    for log in logs_qs.select_related('product', 'source_location', 'destination_location', 'performed_by').order_by('-created_at'):
        row_idx += 1
        dt_str = log.created_at.strftime("%Y-%m-%d %H:%M")
        src_name = log.source_location.name if log.source_location else "—"
        dst_name = log.destination_location.name if log.destination_location else "—"
        actor = log.performed_by.username if log.performed_by else "System"
        ws_logs.append([
            dt_str,
            log.reference_no or "—",
            log.get_log_type_display(),
            log.product.sku,
            log.product.name,
            log.quantity_changed,
            src_name,
            dst_name,
            actor,
            log.reason or "—"
        ])
        r = ws_logs[row_idx]
        for c in r:
            c.font = styles['data_font']
            c.border = styles['border']
        r[5].alignment = Alignment(horizontal="right")
        if row_idx % 2 == 0:
            for c in r:
                c.fill = styles['alt_fill']

    _autofit_columns(ws_logs)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_reports_pdf(start_date, end_date, total_stock_in, total_stock_out, low_stock_list, out_of_stock_list, logs_qs):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1A2332')
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#555555')
    )
    section_style = ParagraphStyle(
        'SectionHeader',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#0D6EFD'),
        spaceBefore=14,
        spaceAfter=6
    )
    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10
    )
    cell_bold = ParagraphStyle(
        'TableCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10
    )

    story = []

    # Title Banner
    story.append(Paragraph("InvenStock — Inventory Analytics Report", title_style))
    story.append(Paragraph(
        f"Period: <b>{start_date}</b> to <b>{end_date}</b> &nbsp;|&nbsp; Generated: {timezone.now().strftime('%Y-%m-%d %H:%M')}",
        subtitle_style
    ))
    story.append(Spacer(1, 14))

    # KPI Summary Table
    story.append(Paragraph("1. Executive Summary", section_style))
    kpi_data = [
        [Paragraph("Metric", cell_bold), Paragraph("Units / Count", cell_bold), Paragraph("Description", cell_bold)],
        [Paragraph("Total Stock In", cell_style), Paragraph(f"+{total_stock_in}", cell_bold), Paragraph("Units added to inventory in period", cell_style)],
        [Paragraph("Total Stock Out", cell_style), Paragraph(f"-{total_stock_out}", cell_bold), Paragraph("Units checked out / sold in period", cell_style)],
        [Paragraph("Net Inventory Change", cell_style), Paragraph(f"{total_stock_in - total_stock_out:+d}", cell_bold), Paragraph("Inflow minus outflow", cell_style)],
        [Paragraph("Low Stock Alert Items", cell_style), Paragraph(f"{len(low_stock_list)}", cell_bold), Paragraph("Items at or below safety threshold", cell_style)],
        [Paragraph("Out of Stock Items", cell_style), Paragraph(f"{len(out_of_stock_list)}", cell_bold), Paragraph("Items completely depleted", cell_style)],
    ]
    kpi_table = Table(kpi_data, colWidths=[2.2 * inch, 1.3 * inch, 3.5 * inch])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A2332')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D9D9D9')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 14))

    # Low Stock Items Table
    story.append(Paragraph("2. Low Stock Alerts", section_style))
    if low_stock_list:
        low_data = [[
            Paragraph("SKU", cell_bold),
            Paragraph("Product Name", cell_bold),
            Paragraph("Current", cell_bold),
            Paragraph("Min Alert", cell_bold),
            Paragraph("Status", cell_bold)
        ]]
        for p in low_stock_list[:25]:
            low_data.append([
                Paragraph(p.sku, cell_style),
                Paragraph(p.name, cell_style),
                Paragraph(str(p.total_stock), cell_bold),
                Paragraph(str(p.low_stock_alert), cell_style),
                Paragraph("Low Stock", cell_style),
            ])
        low_table = Table(low_data, colWidths=[1.3 * inch, 3.2 * inch, 0.8 * inch, 0.8 * inch, 0.9 * inch])
        low_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E65100')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D9D9D9')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FFF8E1')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(low_table)
    else:
        story.append(Paragraph("<i>All products are sufficiently stocked.</i>", subtitle_style))
    story.append(Spacer(1, 14))

    # Out of Stock Items Table
    story.append(Paragraph("3. Out of Stock Products", section_style))
    if out_of_stock_list:
        out_data = [[
            Paragraph("SKU", cell_bold),
            Paragraph("Product Name", cell_bold),
            Paragraph("Category", cell_bold),
            Paragraph("Supplier", cell_bold)
        ]]
        for p in out_of_stock_list[:25]:
            cat_name = p.category.name if p.category else "—"
            sup_name = p.supplier.name if p.supplier else "—"
            out_data.append([
                Paragraph(p.sku, cell_style),
                Paragraph(p.name, cell_style),
                Paragraph(cat_name, cell_style),
                Paragraph(sup_name, cell_style),
            ])
        out_table = Table(out_data, colWidths=[1.3 * inch, 3.1 * inch, 1.3 * inch, 1.3 * inch])
        out_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#C5221F')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D9D9D9')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FFEBEE')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(out_table)
    else:
        story.append(Paragraph("<i>No products are currently out of stock.</i>", subtitle_style))

    doc.build(story)
    buffer.seek(0)
    return buffer


def generate_reports_csv(start_date, end_date, total_stock_in, total_stock_out, low_stock_list, out_of_stock_list, logs_qs=None):
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    # --- SECTION 1: Summary ---
    writer.writerow(["InvenStock — Inventory Analytics Report"])
    writer.writerow([f"Period: {start_date} to {end_date} | Generated: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    writer.writerow([])
    writer.writerow(["Key Metric", "Value", "Notes"])
    summary_rows = [
        ("Total Stock Received (Stock In)", total_stock_in, "Units taken into inventory during selected period"),
        ("Total Stock Dispatched (Stock Out)", total_stock_out, "Units checked out / sold during selected period"),
        ("Net Stock Movement", total_stock_in - total_stock_out, "Difference between inbound and outbound units"),
        ("Low Stock Products Count", len(low_stock_list), "Products at or below their alert threshold"),
        ("Out of Stock Products Count", len(out_of_stock_list), "Products currently having zero inventory"),
    ]
    for item, val, note in summary_rows:
        writer.writerow([item, val, note])
    writer.writerow([])

    # --- SECTION 2: Low Stock Alerts ---
    writer.writerow(["Low Stock Alerts"])
    writer.writerow([f"Products requiring replenishment as of {timezone.now().strftime('%Y-%m-%d')}"])
    writer.writerow(["SKU", "Product Name", "Category", "Supplier", "Current Stock", "Min Alert Level", "Deficit Units"])
    for p in low_stock_list:
        tot = p.total_stock
        deficit = max(0, p.low_stock_alert - tot)
        cat_name = p.category.name if p.category else "—"
        sup_name = p.supplier.name if p.supplier else "—"
        writer.writerow([p.sku, p.name, cat_name, sup_name, tot, p.low_stock_alert, deficit])
    writer.writerow([])

    # --- SECTION 3: Out of Stock ---
    writer.writerow(["Out of Stock Products"])
    writer.writerow([f"Zero-quantity products as of {timezone.now().strftime('%Y-%m-%d')}"])
    writer.writerow(["SKU", "Product Name", "Category", "Supplier", "Retail Price", "Cost Price", "Status"])
    for p in out_of_stock_list:
        cat_name = p.category.name if p.category else "—"
        sup_name = p.supplier.name if p.supplier else "—"
        writer.writerow([
            p.sku,
            p.name,
            cat_name,
            sup_name,
            f"${float(p.retail_price):.2f}",
            f"${float(p.cost_price):.2f}",
            p.get_status_display()
        ])
    writer.writerow([])

    # --- SECTION 4: Stock Transactions Audit ---
    writer.writerow(["Inventory Movements Audit Log"])
    writer.writerow([f"Transactions recorded between {start_date} and {end_date}"])
    headers_logs = [
        "Date / Time", "Reference", "Type", "Product SKU", "Product Name",
        "Qty Changed", "From Location", "To Location", "Performed By", "Reason"
    ]
    writer.writerow(headers_logs)
    if logs_qs is not None:
        qs = logs_qs.select_related('product', 'source_location', 'destination_location', 'performed_by').order_by('-created_at') if hasattr(logs_qs, 'select_related') else logs_qs
        for log in qs:
            dt_str = log.created_at.strftime("%Y-%m-%d %H:%M")
            src_name = log.source_location.name if log.source_location else "—"
            dst_name = log.destination_location.name if log.destination_location else "—"
            actor = log.performed_by.username if log.performed_by else "System"
            writer.writerow([
                dt_str,
                log.reference_no or "—",
                log.get_log_type_display(),
                log.product.sku,
                log.product.name,
                log.quantity_changed,
                src_name,
                dst_name,
                actor,
                log.reason or "—"
            ])

    return buffer.getvalue()


# ─────────────────────────────────────────────
#  Products Inventory Exports
# ─────────────────────────────────────────────

def generate_products_excel(products_qs):
    wb = openpyxl.Workbook()
    styles = _get_styles()
    ws = wb.active
    ws.title = "Products Inventory"
    ws.views.sheetView[0].showGridLines = True

    ws["A1"] = "InvenStock — Products Catalog & Stock Levels"
    ws["A1"].font = styles['title_font']
    ws["A2"] = f"Exported: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')} | Total items: {products_qs.count()}"
    ws["A2"].font = styles['subtitle_font']

    ws.append([])
    headers = [
        "SKU", "Product Name", "Category", "Supplier",
        "Retail Price ($)", "Cost Price ($)",
        "Current Stock", "Low Stock Alert", "Stock Status", "Status"
    ]
    ws.append(headers)
    for cell in ws[4]:
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 4
    for p in products_qs:
        row_idx += 1
        cat_name = p.category.name if p.category else "—"
        sup_name = p.supplier.name if p.supplier else "—"
        tot = p.total_stock
        stock_status = p.stock_status.replace('_', ' ').title()

        ws.append([
            p.sku,
            p.name,
            cat_name,
            sup_name,
            float(p.retail_price),
            float(p.cost_price),
            tot,
            p.low_stock_alert,
            stock_status,
            p.get_status_display()
        ])
        r = ws[row_idx]
        for c in r:
            c.font = styles['data_font']
            c.border = styles['border']
        r[4].number_format = '$#,##0.00'
        r[5].number_format = '$#,##0.00'
        r[6].alignment = Alignment(horizontal="right")
        r[7].alignment = Alignment(horizontal="right")
        if row_idx % 2 == 0:
            for c in r:
                c.fill = styles['alt_fill']

    _autofit_columns(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_products_csv(products_qs):
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    # Title block (matching Excel export format)
    writer.writerow(["InvenStock — Products Catalog & Stock Levels"])
    writer.writerow([f"Exported: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')} | Total items: {products_qs.count()}"])
    writer.writerow([])  # blank spacer row

    headers = [
        "SKU", "Product Name", "Category", "Supplier",
        "Retail Price ($)", "Cost Price ($)", "Current Stock",
        "Low Stock Alert", "Stock Status", "Status"
    ]
    writer.writerow(headers)

    for p in products_qs:
        cat_name = p.category.name if p.category else "—"
        sup_name = p.supplier.name if p.supplier else "—"
        tot = p.total_stock
        stock_status = p.stock_status.replace('_', ' ').title()
        writer.writerow([
            p.sku,
            p.name,
            cat_name,
            sup_name,
            f"{float(p.retail_price):.2f}",
            f"{float(p.cost_price):.2f}",
            tot,
            p.low_stock_alert,
            stock_status,
            p.get_status_display()
        ])

    return buffer.getvalue()
